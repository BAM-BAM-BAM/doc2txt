"""Tests for folder watcher (doc2txt_watcher.py).

FGT categories:
  INV-*  : Watcher invariants
  BOUND-*: Edge case handling
  INT-*  : Integration tests
"""

import sys
import time
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from doc2txt_watcher import (
    FileReadinessChecker,
    ProcessedTracker,
    FolderWatcher,
    WatchConfig,
    _build_arg_parser,
    main,
)


class TestInvFileReadiness:
    """INV/BOUND tests for FileReadinessChecker."""

    @pytest.fixture
    def checker(self):
        return FileReadinessChecker(cooldown_minutes=10, min_file_size=100)

    def test_inv_nonexistent_file_not_ready(self, checker, tmp_path):
        """Non-existent file is not ready."""
        fake = tmp_path / "nonexistent.pdf"
        ready, reason = checker.is_ready(fake)
        assert not ready
        assert "does not exist" in reason

    def test_inv_directory_not_ready(self, checker, tmp_path):
        """Directory is not ready (only files)."""
        d = tmp_path / "subdir"
        d.mkdir()
        ready, reason = checker.is_ready(d)
        assert not ready
        assert "not a regular file" in reason

    def test_inv_small_file_not_ready(self, checker, tmp_path):
        """Files smaller than min_file_size are not ready (placeholders)."""
        small = tmp_path / "tiny.pdf"
        small.write_bytes(b"x" * 50)
        # Set mtime to long ago so cooldown isn't the reason
        import os
        old_time = time.time() - 3600
        os.utime(small, (old_time, old_time))
        ready, reason = checker.is_ready(small)
        assert not ready
        assert "too small" in reason

    def test_inv_recent_file_not_ready(self, checker, tmp_path):
        """File modified less than cooldown_minutes ago is not ready."""
        recent = tmp_path / "recent.pdf"
        recent.write_bytes(b"x" * 200)
        # File was just created, so mtime is now
        ready, reason = checker.is_ready(recent)
        assert not ready
        assert "modified" in reason

    def test_inv_old_file_is_ready(self, tmp_path):
        """File older than cooldown is ready."""
        checker = FileReadinessChecker(cooldown_minutes=0, min_file_size=10)
        old = tmp_path / "old.pdf"
        old.write_bytes(b"x" * 200)
        import os
        old_time = time.time() - 3600
        os.utime(old, (old_time, old_time))
        ready, reason = checker.is_ready(old)
        assert ready
        assert reason == "ready"

    def test_bound_libreoffice_lock_blocks(self, checker, tmp_path):
        """LibreOffice lock file makes document not ready."""
        doc = tmp_path / "document.docx"
        doc.write_bytes(b"x" * 200)
        import os
        old_time = time.time() - 3600
        os.utime(doc, (old_time, old_time))

        # Create LibreOffice lock
        lock = tmp_path / ".~lock.document.docx#"
        lock.write_text("lock")

        ready, reason = checker.is_ready(doc)
        assert not ready
        assert "LibreOffice lock" in reason

    def test_bound_office_lock_blocks(self, checker, tmp_path):
        """Microsoft Office lock file makes document not ready."""
        doc = tmp_path / "report.docx"
        doc.write_bytes(b"x" * 200)
        import os
        old_time = time.time() - 3600
        os.utime(doc, (old_time, old_time))

        # Create Office lock
        lock = tmp_path / "~$report.docx"
        lock.write_text("lock")

        ready, reason = checker.is_ready(doc)
        assert not ready
        assert "Office lock" in reason

    def test_bound_vim_swap_blocks(self, checker, tmp_path):
        """Vim swap file makes document not ready."""
        doc = tmp_path / "notes.md"
        doc.write_bytes(b"x" * 200)
        import os
        old_time = time.time() - 3600
        os.utime(doc, (old_time, old_time))

        # Create vim swap
        swap = tmp_path / ".notes.md.swp"
        swap.write_text("swap")

        ready, reason = checker.is_ready(doc)
        assert not ready
        assert "Vim swap" in reason

    def test_bound_temp_file_not_ready(self, checker, tmp_path):
        """Cloud sync temp files are not ready."""
        tmp_file = tmp_path / "download.tmp"
        tmp_file.write_bytes(b"x" * 200)
        import os
        old_time = time.time() - 3600
        os.utime(tmp_file, (old_time, old_time))

        ready, reason = checker.is_ready(tmp_file)
        assert not ready
        assert "temp pattern" in reason


class TestProcessedTracker:
    """Tests for ProcessedTracker."""

    @pytest.fixture
    def tracker(self, tmp_path):
        db_path = tmp_path / "test_watcher.db"
        t = ProcessedTracker(db_path)
        yield t
        t.close()

    def test_inv_new_file_needs_processing(self, tracker, tmp_path):
        """A never-seen file needs processing."""
        doc = tmp_path / "new.pdf"
        doc.write_bytes(b"PDF content")
        assert tracker.needs_processing(doc)

    def test_inv_processed_file_skipped(self, tracker, tmp_path):
        """A processed file with unchanged mtime is skipped."""
        doc = tmp_path / "done.pdf"
        doc.write_bytes(b"PDF content")
        tracker.mark_processed(doc)
        assert not tracker.needs_processing(doc)

    def test_inv_modified_file_needs_reprocessing(self, tracker, tmp_path):
        """A processed file with changed mtime needs reprocessing."""
        doc = tmp_path / "modified.pdf"
        doc.write_bytes(b"PDF content v1")
        tracker.mark_processed(doc)

        # Modify the file (changes mtime and size)
        time.sleep(0.1)
        doc.write_bytes(b"PDF content v2 with more data")
        assert tracker.needs_processing(doc)

    def test_int_stats(self, tracker, tmp_path):
        """Stats reflect tracked files."""
        doc = tmp_path / "file.pdf"
        doc.write_bytes(b"content")
        tracker.mark_processed(doc)
        stats = tracker.get_stats()
        assert stats["total_tracked"] == 1


class TestFolderWatcher:
    """Integration tests for FolderWatcher."""

    def _make_old_file(self, path: Path, content: bytes = b"x" * 200):
        """Create a file with mtime in the past."""
        import os
        path.write_bytes(content)
        old_time = time.time() - 3600
        os.utime(path, (old_time, old_time))

    def test_int_scan_finds_ready_files(self, tmp_path):
        """scan_once finds files that are old enough and supported."""
        self._make_old_file(tmp_path / "ready.pdf")
        self._make_old_file(tmp_path / "ready.docx")
        (tmp_path / "too_new.pdf").write_bytes(b"x" * 200)  # mtime = now
        self._make_old_file(tmp_path / "ignored.txt")  # wrong extension

        config = WatchConfig(
            watch_dirs=[tmp_path],
            cooldown_minutes=5,  # 5 min cooldown so "too_new" is blocked
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()

        ready = [p.name for p, r in results if r == "ready"]
        pending = [p.name for p, r in results if r != "ready"]
        assert "ready.pdf" in ready
        assert "ready.docx" in ready
        assert "too_new.pdf" in pending  # Too recently modified
        assert "ignored.txt" not in ready and "ignored.txt" not in pending  # Wrong ext

    def test_int_scan_skips_processed(self, tmp_path):
        """scan_once skips already-processed files."""
        self._make_old_file(tmp_path / "already.pdf")

        config = WatchConfig(
            watch_dirs=[tmp_path],
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)

        # First scan: file is new
        results1 = watcher.scan_once()
        assert len(results1) == 1

        # Mark as processed
        watcher.tracker.mark_processed(tmp_path / "already.pdf")

        # Second scan: file is skipped
        results2 = watcher.scan_once()
        assert len(results2) == 0

        watcher.close()

    def test_int_scan_recursive(self, tmp_path):
        """scan_once finds files in subdirectories when recursive."""
        sub = tmp_path / "subdir"
        sub.mkdir()
        self._make_old_file(sub / "nested.pdf")

        config = WatchConfig(
            watch_dirs=[tmp_path],
            recursive=True,
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()

        ready = [p.name for p, r in results if r == "ready"]
        assert "nested.pdf" in ready

    def test_bound_unreachable_watch_dir_skipped(self, tmp_path, monkeypatch):
        """BUG-004: OSError on one watch dir (dead mount) must not abort the scan.

        Simulates an unmounted drive where exists() raises ENODEV instead of
        returning False. The other watch dir must still be scanned.
        """
        good_dir = tmp_path / "good"
        good_dir.mkdir()
        self._make_old_file(good_dir / "reachable.pdf")
        dead_mount = tmp_path / "dead_mount"

        real_exists = Path.exists

        def fake_exists(self, **kwargs):
            if self == dead_mount:
                raise OSError(19, "No such device", str(self))
            return real_exists(self, **kwargs)

        monkeypatch.setattr(Path, "exists", fake_exists)

        config = WatchConfig(
            watch_dirs=[dead_mount, good_dir],
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()

        ready = [p.name for p, r in results if r == "ready"]
        assert "reachable.pdf" in ready

    def test_bound_io_error_on_one_file_skips_file(self, tmp_path, monkeypatch):
        """BUG-004: OSError on one file mid-scan must skip that file only.

        Simulates a cloud-sync file where stat() raises EIO. The remaining
        files in the same dir must still be scanned.
        """
        self._make_old_file(tmp_path / "healthy.pdf")
        bad_file = tmp_path / "stalled_sync.pdf"
        self._make_old_file(bad_file)

        real_stat = Path.stat

        def fake_stat(self, **kwargs):
            if self == bad_file:
                raise OSError(5, "Input/output error", str(self))
            return real_stat(self, **kwargs)

        monkeypatch.setattr(Path, "stat", fake_stat)

        config = WatchConfig(
            watch_dirs=[tmp_path],
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()

        names = [p.name for p, _ in results]
        assert "healthy.pdf" in names
        assert "stalled_sync.pdf" not in names

    def test_bound_unlistable_watch_dir_skipped(self, tmp_path, monkeypatch):
        """BUG-004: OSError while listing one watch dir skips that dir only."""
        good_dir = tmp_path / "good"
        good_dir.mkdir()
        self._make_old_file(good_dir / "reachable.pdf")
        bad_dir = tmp_path / "bad"
        bad_dir.mkdir()

        real_rglob = Path.rglob

        def fake_rglob(self, pattern):
            if self == bad_dir:
                raise OSError(5, "Input/output error", str(self))
            return real_rglob(self, pattern)

        monkeypatch.setattr(Path, "rglob", fake_rglob)

        config = WatchConfig(
            watch_dirs=[bad_dir, good_dir],
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()

        ready = [p.name for p, r in results if r == "ready"]
        assert "reachable.pdf" in ready

    def test_int_scan_nonexistent_dir(self, tmp_path):
        """scan_once handles non-existent watch directory gracefully."""
        config = WatchConfig(
            watch_dirs=[tmp_path / "nonexistent"],
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        results = watcher.scan_once()
        watcher.close()
        assert results == []

    def test_int_process_creates_markdown(self, tmp_path):
        """process_ready_files creates .md output for DOCX files."""
        import docx
        doc = docx.Document()
        doc.add_paragraph("Watched content")
        docx_path = tmp_path / "watched.docx"
        doc.save(str(docx_path))
        self._make_old_file(docx_path, docx_path.read_bytes())

        config = WatchConfig(
            watch_dirs=[tmp_path],
            cooldown_minutes=0,
            min_file_size=10,
            db_path=tmp_path / "watcher.db",
        )
        watcher = FolderWatcher(config)
        scan_results = watcher.scan_once()

        stats = watcher.process_ready_files(scan_results)
        watcher.close()

        assert stats["processed"] == 1
        md_path = tmp_path / "watched.md"
        assert md_path.exists()
        assert "Watched content" in md_path.read_text()


class TestCli:
    """Tests for the CLI entrypoint (_build_arg_parser / main)."""

    def test_schema_once_flag(self):
        args = _build_arg_parser().parse_args(["--watch-dir", "/tmp/a", "--once"])
        assert args.once is True

    def test_schema_once_defaults_false(self):
        args = _build_arg_parser().parse_args(["--watch-dir", "/tmp/a"])
        assert args.once is False

    def test_schema_multiple_watch_dirs_accumulate(self):
        args = _build_arg_parser().parse_args(
            ["--watch-dir", "/tmp/a", "-d", "/tmp/b"]
        )
        assert args.watch_dir == [Path("/tmp/a"), Path("/tmp/b")]

    def test_schema_watch_dir_required(self, capsys):
        with pytest.raises(SystemExit):
            _build_arg_parser().parse_args(["--once"])

    def test_schema_cooldown_parsed_as_int(self):
        args = _build_arg_parser().parse_args(
            ["--watch-dir", "/tmp/a", "--cooldown", "45"]
        )
        assert args.cooldown == 45

    def test_schema_defaults(self):
        args = _build_arg_parser().parse_args(["--watch-dir", "/tmp/a"])
        assert args.cooldown == 10
        assert args.poll_interval == 30
        assert args.dry_run is False
        assert args.verbose is False

    def test_bound_unsupported_format_rejected(self, tmp_path):
        with pytest.raises(SystemExit, match="Unsupported format"):
            main(["--watch-dir", str(tmp_path), "--formats", "exe",
                  "--db-path", str(tmp_path / "w.db"), "--once"])

    def _capture_watcher(self, monkeypatch, scan_results):
        """Stub FolderWatcher's scan/run/process methods, recording calls."""
        calls = {"scan": 0, "run": 0, "process": 0, "config": None}

        real_init = FolderWatcher.__init__

        def fake_init(self, config):
            real_init(self, config)
            calls["config"] = config

        def fake_scan(self):
            calls["scan"] += 1
            return scan_results

        def fake_run(self):
            calls["run"] += 1

        def fake_process(self, results):
            calls["process"] += 1
            return {"processed": len(results), "skipped": 0, "failed": 0}

        monkeypatch.setattr(FolderWatcher, "__init__", fake_init)
        monkeypatch.setattr(FolderWatcher, "scan_once", fake_scan)
        monkeypatch.setattr(FolderWatcher, "run", fake_run)
        monkeypatch.setattr(FolderWatcher, "process_ready_files", fake_process)
        return calls

    def test_int_main_once_dispatches_single_scan(self, tmp_path, monkeypatch):
        """--once runs exactly one scan and never enters the watch loop."""
        calls = self._capture_watcher(monkeypatch, scan_results=[])
        main(["--watch-dir", str(tmp_path), "--once",
              "--db-path", str(tmp_path / "w.db")])
        assert calls["scan"] == 1
        assert calls["run"] == 0
        assert calls["process"] == 0  # no ready files -> no processing

    def test_int_main_once_processes_ready_files(self, tmp_path, monkeypatch):
        """--once with ready scan results dispatches to process_ready_files."""
        ready = [(tmp_path / "a.pdf", "ready")]
        calls = self._capture_watcher(monkeypatch, scan_results=ready)
        main(["--watch-dir", str(tmp_path), "--once",
              "--db-path", str(tmp_path / "w.db")])
        assert calls["scan"] == 1
        assert calls["process"] == 1

    def test_int_main_without_once_enters_watch_loop(self, tmp_path, monkeypatch):
        """Without --once, main dispatches to the polling loop."""
        calls = self._capture_watcher(monkeypatch, scan_results=[])
        main(["--watch-dir", str(tmp_path),
              "--db-path", str(tmp_path / "w.db")])
        assert calls["run"] == 1
        assert calls["scan"] == 0

    def test_int_main_builds_config_from_args(self, tmp_path, monkeypatch):
        """CLI args land in the WatchConfig the watcher is constructed with."""
        calls = self._capture_watcher(monkeypatch, scan_results=[])
        main(["--watch-dir", str(tmp_path), "--once", "--cooldown", "7",
              "--formats", "pdf,docx", "--dry-run",
              "--db-path", str(tmp_path / "w.db")])
        config = calls["config"]
        assert config.watch_dirs == [tmp_path]
        assert config.cooldown_minutes == 7
        assert config.formats == {".pdf", ".docx"}
        assert config.dry_run is True
